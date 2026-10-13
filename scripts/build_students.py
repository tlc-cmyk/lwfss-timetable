#!/usr/bin/env python3
"""Build students.js / students.json from class-register.xlsx.

Source sheets: 1A–6D with columns No, English Name, Alias, 中文姓名, Sex, House.
Public output: Chinese display name + class + class number + Sex + House (no emails).
"""
from __future__ import annotations

import datetime
import json
import re
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = Path(__file__).resolve().parent / "class-register.xlsx"
CLASS_RE = re.compile(r"^[1-6][A-Da-d]$")


def cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def find_header_row(ws) -> int | None:
    for i, row in enumerate(ws.iter_rows(max_row=30, values_only=True), 1):
        if not row:
            continue
        a = cell_str(row[0]) if len(row) > 0 else ""
        d = cell_str(row[3]) if len(row) > 3 else ""
        if a == "No" and ("中文" in d or d == "中文姓名"):
            return i
    return None


def parse_xlsx(xlsx_path: Path) -> dict:
    try:
        import openpyxl
    except ImportError as e:
        raise SystemExit(
            "openpyxl required. Try: source /tmp/xlsvenv/bin/activate"
        ) from e

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    students: list[dict] = []
    eng_fallback = 0

    for sheet_name in wb.sheetnames:
        if not CLASS_RE.match(sheet_name):
            continue
        cls = sheet_name.upper()
        ws = wb[sheet_name]
        header_i = find_header_row(ws)
        if header_i is None:
            print(f"warn: no header on sheet {sheet_name}", file=sys.stderr)
            continue
        for row in ws.iter_rows(min_row=header_i + 1, values_only=True):
            if not row:
                continue
            no_raw = row[0] if len(row) > 0 else None
            eng = cell_str(row[1]) if len(row) > 1 else ""
            chi = cell_str(row[3]) if len(row) > 3 else ""
            no_s = cell_str(no_raw)
            if not no_s:
                continue
            try:
                no_int = int(float(no_s))
            except (TypeError, ValueError):
                continue
            if no_int < 1 or no_int > 99:
                continue
            no = f"{no_int:02d}"
            if not chi:
                if not eng:
                    continue
                chi = eng.title() if eng.isupper() else eng
                eng_fallback += 1
            sex = cell_str(row[4]) if len(row) > 4 else ""
            house = cell_str(row[5]) if len(row) > 5 else ""
            sex = sex.upper()[:1] if sex else ""
            if sex not in ("M", "F"):
                sex = ""
            house = house.upper()[:1] if house else ""
            if house not in ("S", "W", "E", "N"):
                house = ""
            rec = {"class": cls, "no": no, "chi": chi}
            if sex:
                rec["sex"] = sex
            if house:
                rec["house"] = house
            students.append(rec)

    students.sort(key=lambda s: (s["class"], s["no"]))
    by_class = Counter(s["class"] for s in students)
    meta = {
        "school": "世界龍岡學校劉皇發中學",
        "source": "Class Register Excel (sheets 1A–6D)",
        "generatedAt": datetime.datetime.now(
            datetime.timezone(datetime.timedelta(hours=8))
        ).strftime("%Y-%m-%dT%H:%M:%S+08:00"),
        "timezone": "Asia/Hong_Kong",
        "studentCount": len(students),
        "classCount": len(by_class),
        "classes": dict(sorted(by_class.items())),
        "fields": ["class", "no", "chi", "sex", "house"],
        "privacy": "Chinese display name + class + class number + Sex + House; no emails",
        "note": (
            "From class register: No + 中文姓名 + Sex + House; "
            "English Name used as display only when 中文姓名 blank"
            + (f" ({eng_fallback} fallbacks)" if eng_fallback else "")
        ),
    }
    return {"meta": meta, "students": students}


def write_outputs(payload: dict) -> None:
    (ROOT / "students.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    js = (
        "// Auto-generated student roster for 模擬遊戲 — do not edit by hand.\n"
        "// Rebuild: python3 scripts/build_students.py [class-register.xlsx]\n"
        "// Privacy: Chinese names + class + Sex + House (no emails).\n"
        "window.STUDENTS_DATA = "
        + json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
        + ";\n"
    )
    (ROOT / "students.js").write_text(js, encoding="utf-8")


def main() -> int:
    xlsx_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not xlsx_path.is_file():
        print(f"XLSX not found: {xlsx_path}", file=sys.stderr)
        return 1
    payload = parse_xlsx(xlsx_path)
    write_outputs(payload)
    m = payload["meta"]
    print(
        f"Wrote {m['studentCount']} students / {m['classCount']} classes "
        f"→ students.js / students.json"
    )
    print("classes:", m["classes"])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
